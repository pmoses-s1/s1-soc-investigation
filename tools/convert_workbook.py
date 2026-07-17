#!/usr/bin/env python3
"""
Convert a DFIR investigation workbook (.xlsx) into engine catalog YAML files.

    python tools/convert_workbook.py <workbook.xlsx> [output_dir=catalogs]

Emits one catalog per DFIR domain (dfir_identity_access, dfir_endpoint, ...) plus a
master dfir_insider_threat_full. Maps workbook placeholders (%EMAIL% -> {{entity}},
%HOSTNAME%/%AGENTUUID%/%IP%/%USERNAME%/%SF_USER_ID% -> template vars), parameterizes
config datatable names to optional {{dt_*}} variables, expands the non-SDL
`any(f1,...) contains:anycase("x")` selector into an explicit OR, unescapes workbook
paste-escapes, normalizes bare `contains` to `contains:anycase`, and drops non-query
reference rows. Requires openpyxl and PyYAML.
"""
import re, sys
from pathlib import Path
import openpyxl, yaml

SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else None
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("catalogs")
if not SRC or not SRC.is_file():
    sys.exit("usage: convert_workbook.py <workbook.xlsx> [output_dir]")

TAB_DOMAIN = {
 "🔍 Coverage":"identity_access","🏠 User Lookup":"identity_access","🔑 Okta":"identity_access",
 "☁️ Azure":"identity_access","👤 HR Identity":"identity_access","🔐 Secrets":"identity_access",
 "🖥️ Endpoint":"endpoint","🔒 Op Lockdown":"endpoint",
 "📁 Google WS":"collab_storage","💬 Slack":"collab_storage","📦 Storage Exfil":"collab_storage","🧩 Atlassian":"collab_storage",
 "🌐 ZIA":"web_network","🌐 ZIA Insider":"web_network","🌐 Network DNS":"web_network",
 "☁️ AWS Cloud":"cloud",
 "💼 Salesforce":"saas_apps","🐙 GitHub":"saas_apps","🧠 SaaS Misc":"saas_apps","🤖 Prompt Sec":"saas_apps","✈️ Travel":"saas_apps",
 "🛡️ DLP Triage":"exfil_dlp","🧲 ZIA DLP Endpoint":"exfil_dlp",
 "🔗 Endpoint Correlation":"correlation","⚡ Quick Pivots":"correlation",
}
DOMAIN_NAME = {
 "identity_access":"DFIR: Identity & access","endpoint":"DFIR: Endpoint","collab_storage":"DFIR: Collaboration & storage",
 "web_network":"DFIR: Web & network","cloud":"DFIR: Cloud","saas_apps":"DFIR: SaaS & apps",
 "exfil_dlp":"DFIR: Exfil & DLP","correlation":"DFIR: Cross-source correlation",
}
# Canonical placeholder -> template variable. Covers both %X% and <X> delimiters.
# Anything not listed here still gets parameterized generically (see mapph) so no
# placeholder is ever left as a literal string in a query.
PLACE = [("%EMAIL%","{{entity}}"),("%USERNAME%","{{username}}"),("%AGENTUUID%","{{agent_uuid}}"),
 ("%HOSTNAME%","{{hostname}}"),("%IP%","{{ip}}"),("%SF_USER_ID%","{{sf_user_id}}"),
 ("%SESSION%","{{session}}"),("%SESSION_KEY%","{{session}}"),("%LOGIN_KEY%","{{login_key}}"),
 ("%DOMAIN%","{{domain}}"),("%APP_NAME%","{{app_name}}"),("%FILE_OR_TITLE%","{{file_or_title}}"),
 ("%USER%","{{username}}"),("%HOST%","{{hostname}}"),
 ("<IP>","{{ip}}"),("<SESSION>","{{session}}"),("<SESSION_KEY>","{{session}}"),
 ("<HOST>","{{hostname}}"),("<HOSTNAME>","{{hostname}}"),("<USER>","{{username}}"),
 ("<USERNAME>","{{username}}"),("<AGENTUUID>","{{agent_uuid}}"),("<EMAIL>","{{entity}}"),
 ("<LOGIN_KEY>","{{login_key}}"),("<DOMAIN>","{{domain}}"),("<APP_NAME>","{{app_name}}"),
 ("<FILE_OR_TITLE>","{{file_or_title}}"),("<SF_USER_ID>","{{sf_user_id}}")]
DT_VARS = set()
SRC_VARS = set()      # serverHost source names turned into overridable {{src_*|default}}
GENERIC_VARS = set()  # any other placeholder we auto-slugged, for reporting

def _slug_var(tok):
    return re.sub(r'[^a-z0-9]+', '_', tok.lower()).strip('_')

def mapph(s):
    for a,b in PLACE: s=s.replace(a,b)
    # Generic fallback: convert any remaining %X% or <X> delimited placeholder to a
    # {{slug}} variable so nothing is left literal (which would match no rows).
    def genrepl(m):
        var=_slug_var(m.group(1)); GENERIC_VARS.add(var)
        return "{{"+var+"}}"
    s=re.sub(r'%([A-Za-z][A-Za-z0-9_]*)%', genrepl, s)
    s=re.sub(r'<([A-Za-z][A-Za-z0-9_]*)>', genrepl, s)
    return s

def param_serverhost(s):
    """serverHost='zia' -> serverHost='{{src_zia|zia}}' (value kept as default,
    so it runs out of the box but is overridable per tenant). Empty is left as-is."""
    def repl(m):
        q, val = m.group(1), m.group(2)
        if not val.strip():
            return m.group(0)
        var="src_"+_slug_var(val); SRC_VARS.add((var, val))
        return f"serverHost={q}{{{{{var}|{val}}}}}{q}"
    return re.sub(r"serverHost\s*=\s*(['\"])(.*?)\1", repl, s)

def expand_any(s):
    pat = re.compile(r'any\(\s*(.*?)\)\s*contains:(anycase|matchcase)\(\s*("(?:[^"]*)"|\'(?:[^\']*)\')\s*\)', re.DOTALL)
    def repl(m):
        fields=[f.strip() for f in m.group(1).split(",") if f.strip()]
        return "(" + " or ".join(f'{f} contains:{m.group(2)}({m.group(3)})' for f in fields) + ")"
    return pat.sub(repl, s)

def param_datatables(s):
    def repl(m):
        slug="dt_"+re.sub(r'[^a-z0-9]+','_',m.group(1).lower()).strip('_'); DT_VARS.add(slug)
        return "config://datatables/{{"+slug+"}}"
    return re.sub(r'config://datatables/([A-Za-z0-9_./-]+)', repl, s)

def normalize(s):
    s=s.replace('\\"','"').replace('\\ ',' ')
    s=re.sub(r"contains\s+'([^']*)'", r'contains:anycase("\1")', s)
    s=re.sub(r'contains\s+"([^"]*)"', r'contains:anycase("\1")', s)
    s=expand_any(s); s=param_datatables(s); s=param_serverhost(s)
    return s

def looks_like_query(pq):
    return ("|" in pq) or ("=" in pq) or bool(re.search(
        r'\b(serverHost|dataSource|filter|group|union|dataset|columns|sort|limit)\b', pq))

class LS(str): pass
yaml.add_representer(LS, lambda d,x: d.represent_scalar('tag:yaml.org,2002:str', x, style='|'))
def slug(s): return re.sub(r'[^a-z0-9]+','-',str(s).strip().lower()).strip('-')

def header_row(rows):
    for i,r in enumerate(rows):
        v=[("" if c is None else str(c)).strip() for c in r]
        if v and v[0]=="ID" and any("PowerQuery" in x for x in v): return i,v
        if v and v[0]=="Pivot" and any("Query Template" in x for x in v): return i,v
    return None,None

def colidx(hdr,*keys):
    for i,h in enumerate(hdr):
        if any(k.lower() in h.lower() for k in keys): return i
    return None

wb=openpyxl.load_workbook(SRC, data_only=True, read_only=True)
domains={d:[] for d in set(TAB_DOMAIN.values())}; master=[]; seen=set(); dropped=0
for ws in wb.worksheets:
    dom=TAB_DOMAIN.get(ws.title)
    if not dom: continue
    rows=[list(r) for r in ws.iter_rows(values_only=True)]
    hi,hdr=header_row(rows)
    if hi is None: continue
    quick=hdr[0]=="Pivot"
    if quick:
        ci=dict(id=None,name=colidx(hdr,"Pivot"),q=colidx(hdr,"When to use"),pq=colidx(hdr,"Query Template"),src=colidx(hdr,"Primary Sources"),piv=None)
    else:
        ci=dict(id=colidx(hdr,"ID"),name=colidx(hdr,"Query Name"),q=colidx(hdr,"Investigative Question","Investigative Question / Use"),
                pq=colidx(hdr,"PowerQuery"),src=colidx(hdr,"Source"),piv=colidx(hdr,"Pivot To if Positive"))
    n=0
    for r in rows[hi+1:]:
        def cell(i):
            if i is None or i>=len(r) or r[i] is None: return ""
            return str(r[i]).strip()
        pq=cell(ci["pq"])
        if not pq or not looks_like_query(pq):
            if pq: dropped+=1
            continue
        rid=cell(ci["id"]) if ci["id"] is not None else f"QP-{n+1:02d}"
        if not rid: continue
        qid=slug(rid)
        if qid in seen: qid=f"{dom[:3]}-{qid}"
        if qid in seen: continue
        seen.add(qid)
        title=cell(ci["name"]) or rid
        parts=[]; q=cell(ci["q"]); piv=cell(ci["piv"]) if ci["piv"] is not None else ""; src=cell(ci["src"])
        if q: parts.append(q)
        if src: parts.append(f"Source: {src}")
        if piv: parts.append(f"Pivot if positive: {piv}")
        notes=" | ".join(parts)
        pqm=normalize(mapph(pq)).replace("\r\n","\n").rstrip()
        e={"id":qid,"title":title[:120]}
        if notes: e["notes"]=notes
        e["pq"]=LS(pqm) if "\n" in pqm else pqm
        domains[dom].append(e); master.append(dict(e)); n+=1

def dump(path,name,qs):
    path.write_text(yaml.dump({"name":name,"queries":qs},sort_keys=False,allow_unicode=True,width=100000))

OUT.mkdir(parents=True, exist_ok=True)
for dom,qs in domains.items():
    if qs: dump(OUT/f"dfir_{dom}.yaml", DOMAIN_NAME[dom], qs)
dump(OUT/"dfir_insider_threat_full.yaml","DFIR: Insider threat (full sweep)", master)
print(f"wrote {sum(1 for qs in domains.values() if qs)} domain catalogs + master ({len(master)} queries, dropped {dropped})")
print("datatable vars:", sorted(DT_VARS))
print("source vars   :", sorted(SRC_VARS))
print("generic vars  :", sorted(GENERIC_VARS))
